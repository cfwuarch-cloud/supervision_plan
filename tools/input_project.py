"""
輸入工程基本資料 — CLI 及 GUI 雙模式

修正歷程：
  v1.0 2026-06-14 OpenCode Assistant — 初版

作者：OpenCode Assistant / cfwuarch
版本：1.0
最後更新：2026-06-14
相依套件：openpyxl>=3.0.0,<4.0.0; python-docx>=1.1.0,<2.0.0
使用方法：
  python -X utf8 tools/input_project.py           # GUI 模式
  python -X utf8 tools/input_project.py --cli     # CLI 模式
  python -X utf8 tools/input_project.py --test    # CLI + 產生 1 頁測試 Word
功能說明：
  提供工程基本資料輸入（CLI/GUI），
  自動從詳細價目表提取主要工程項目，
  資料儲存為 JSON，可產生測試 Word 文件。
參數說明：
  --cli      命令列互動模式
  --test     測試模式（CLI + 產出 1 頁 Word）
  -o OUT     輸出 JSON 路徑（預設：data/project_info.json）
"""
import argparse
import json
import os
import sys
import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(__file__))))
from common.docx_table import add_cell

# ── 預設值 ──
DEFAULT_JSON = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                            'data', 'project_info.json')
PRICE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                          'data', '02_成德-詳細價目表.xlsx')

DEFAULT_DATA = {
    '工程名稱': '臺南市政府社會局委託辦理北區成德公設民營托嬰中心室內裝修統包工程',
    '工程主辦機關': '臺南市政府社會局',
    '統包廠商': '筑墨室內裝修有限公司',
    '監造單位': '吳真福建築師事務所',
    '工程地點': '臺南市北區成德里立體停車場一樓(地號:北區小北段998-58、998-466 地號土地)',
    '預計工期': '200日曆天',
    '總價': '新台幣14,200,000元整',
    '主要工程項目': [],
}


# ── 工具函式 ──
def extract_major_items(price_path):
    """從詳細價目表提取第二層工程項目（壹.一、壹.二、…）"""
    items = []
    if not os.path.exists(price_path):
        return items
    wb = openpyxl.load_workbook(price_path, data_only=True)
    ws = wb.active
    seen = {}
    for row in ws.iter_rows(min_row=1, max_row=600, values_only=True):
        item = str(row[0]).strip() if row[0] else ''
        name = str(row[1]).strip() if row[1] else ''
        if item and item.count('.') == 1 and name and item.startswith('壹.'):
            # 排除明顯非工程項目的收尾項
            if any(k in name for k in ['小計', '合計', '總價',
                                        '管理費', '利潤', '營業稅', '保險費']):
                continue
            seen[item] = name
    for k in seen:
        items.append(f'{k}  {seen[k]}')
    wb.close()
    return items


def load_data(json_path):
    if os.path.exists(json_path):
        with open(json_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return dict(DEFAULT_DATA)


def save_data(data, json_path):
    os.makedirs(os.path.dirname(json_path), exist_ok=True)
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


CN_NUMS = ['一', '二', '三', '四', '五', '六', '七', '八', '九', '十']


def make_test_word(data, output_path):
    """產出 1 頁 Word 確認基本資料"""
    doc = Document()
    body = doc.element.body

    def add_para(text, bold=False, size=14, center=True):
        p = etree.SubElement(body, qn('w:p'))
        pPr = etree.SubElement(p, qn('w:pPr'))
        jc = etree.SubElement(pPr, qn('w:jc'))
        jc.set(qn('w:val'), 'center' if center else 'left')
        r = etree.SubElement(p, qn('w:r'))
        rPr = etree.SubElement(r, qn('w:rPr'))
        rFonts = etree.SubElement(rPr, qn('w:rFonts'))
        rFonts.set(qn('w:ascii'), '標楷體')
        rFonts.set(qn('w:eastAsia'), '標楷體')
        if bold:
            etree.SubElement(rPr, qn('w:b'))
        sz = etree.SubElement(rPr, qn('w:sz'))
        sz.set(qn('w:val'), str(size * 2))
        szCs = etree.SubElement(rPr, qn('w:szCs'))
        szCs.set(qn('w:val'), str(size * 2))
        t = etree.SubElement(r, qn('w:t'))
        t.text = text
        t.set(qn('xml:space'), 'preserve')

    add_para('工程基本資料', bold=True, size=18)
    add_para('', size=10)

    fields = [
        ('工程名稱', '壹'),
        ('工程主辦機關', '貳'),
        ('統包廠商', '參'),
        ('監造單位', '肆'),
        ('工程地點', '伍'),
        ('預計工期', '陸'),
        ('總價', '柒'),
    ]
    for k, seq in fields:
        add_para(f'{seq}、{k}：{data.get(k, "")}', size=12, center=False)

    add_para('捌、主要工程項目：', bold=True, size=12, center=False)

    items = data.get('主要工程項目', [])
    tbl = etree.SubElement(body, qn('w:tbl'))
    tblPr = etree.SubElement(tbl, qn('w:tblPr'))
    tblW = etree.SubElement(tblPr, qn('w:tblW'))
    tblW.set(qn('w:w'), '9000')
    tblW.set(qn('w:type'), 'dxa')
    tblGrid = etree.SubElement(tbl, qn('w:tblGrid'))
    for w in (1500, 7500):
        gc = etree.SubElement(tblGrid, qn('w:gridCol'))
        gc.set(qn('w:w'), str(w))

    tr_h = etree.SubElement(tbl, qn('w:tr'))
    add_cell(tr_h, '項次', 1500, center=True, bold=True)
    add_cell(tr_h, '項目名稱', 7500, bold=True)

    for item in items:
        parts = item.split(None, 1)
        seq = parts[0] if len(parts) > 0 else ''
        name = parts[1] if len(parts) > 1 else item
        tr_d = etree.SubElement(tbl, qn('w:tr'))
        add_cell(tr_d, seq, 1500, center=True)
        add_cell(tr_d, name, 7500)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f'測試 Word 已產出：{output_path}')


# ── CLI 模式 ──
def cli_mode(data, json_path):
    print('=== 工程基本資料輸入（CLI）===')
    field_seqs = [
        ('工程名稱', '壹'),
        ('工程主辦機關', '貳'),
        ('統包廠商', '參'),
        ('監造單位', '肆'),
        ('工程地點', '伍'),
        ('預計工期', '陸'),
        ('總價', '柒'),
    ]
    for k, seq in field_seqs:
        current = data.get(k, '')
        val = input(f'{seq}、{k} [{current}]: ').strip()
        if val:
            data[k] = val

    print('\n捌、主要工程項目')
    major = extract_major_items(PRICE_PATH)
    if not major:
        print('（無法從價目表提取，請手動輸入）')
        custom = input('請輸入主要工程項目（逗號分隔）: ').strip()
        data['主要工程項目'] = [x.strip() for x in custom.split(',') if x.strip()]
    else:
        print('以下為自動提取之主要工程項目：')
        for i, m in enumerate(major, 1):
            print(f'  {i}. {m}')
        data['主要工程項目'] = major

    save_data(data, json_path)
    print(f'\n資料已儲存至：{json_path}')
    return data


# ── GUI 模式 ──
class ProjectInfoGUI:
    def __init__(self, data, json_path):
        self.data = data
        self.json_path = json_path
        self.root = tk.Tk()
        self.root.title('工程基本資料輸入')
        self.root.geometry('700x600')
        self.entries = {}
        self._build()

    def _build(self):
        main = ttk.Frame(self.root, padding=15)
        main.pack(fill=tk.BOTH, expand=True)

        # 文字欄位
        fields = ['工程名稱', '工程主辦機關', '統包廠商', '監造單位',
                  '工程地點', '預計工期', '總價']
        for i, k in enumerate(fields):
            ttk.Label(main, text=k).grid(row=i, column=0, sticky='w', pady=4)
            var = tk.StringVar(value=self.data.get(k, ''))
            width = 60 if k in ('工程名稱', '工程地點') else 40
            entry = ttk.Entry(main, textvariable=var, width=width)
            entry.grid(row=i, column=1, sticky='ew', padx=8, pady=4)
            self.entries[k] = var

        # 主要工程項目
        ttk.Label(main, text='主要工程項目').grid(
            row=len(fields), column=0, sticky='nw', pady=(12, 4))
        self.items_text = tk.Text(main, height=8, width=70)
        self.items_text.grid(row=len(fields), column=1, sticky='ew',
                             padx=8, pady=(12, 4))
        major = extract_major_items(PRICE_PATH)
        if major:
            self.items_text.insert('1.0', '\n'.join(major))
            ttk.Label(main, text='（自動提取自詳細價目表，可編輯）').grid(
                row=len(fields) + 1, column=1, sticky='w', padx=8)
        else:
            ttk.Label(main, text='（無法讀取價目表，請手動輸入）').grid(
                row=len(fields) + 1, column=1, sticky='w', padx=8)

        # 按鈕
        btn_frame = ttk.Frame(main)
        btn_frame.grid(row=len(fields) + 2, column=0, columnspan=2, pady=20)
        ttk.Button(btn_frame, text='儲存', command=self._save).pack(
            side=tk.LEFT, padx=6)
        ttk.Button(btn_frame, text='儲存並產出測試 Word',
                   command=self._save_and_test).pack(side=tk.LEFT, padx=6)
        ttk.Button(btn_frame, text='離開', command=self.root.destroy).pack(
            side=tk.LEFT, padx=6)

        main.columnconfigure(1, weight=1)

    def _collect(self):
        for k, var in self.entries.items():
            self.data[k] = var.get().strip()
        raw = self.items_text.get('1.0', tk.END).strip()
        self.data['主要工程項目'] = [
            x.strip() for x in raw.split('\n') if x.strip()]

    def _save(self):
        self._collect()
        save_data(self.data, self.json_path)
        messagebox.showinfo('完成', f'資料已儲存至：\n{self.json_path}')

    def _save_and_test(self):
        self._collect()
        save_data(self.data, self.json_path)
        out_word = os.path.join(
            os.path.dirname(self.json_path), '..',
            'output', 'project_info_test.docx')
        out_word = os.path.normpath(out_word)
        make_test_word(self.data, out_word)
        messagebox.showinfo('完成', f'資料已儲存，測試 Word：\n{out_word}')

    def run(self):
        self.root.mainloop()


# ── 主程式 ──
def main():
    parser = argparse.ArgumentParser(
        description='輸入工程基本資料（CLI / GUI 雙模式）')
    parser.add_argument('--cli', action='store_true',
                        help='命令列互動模式（預設為 GUI）')
    parser.add_argument('--test', action='store_true',
                        help='測試模式（CLI + 產出 1 頁 Word）')
    parser.add_argument('-o', default=DEFAULT_JSON,
                        help=f'輸出 JSON 路徑（預設：{DEFAULT_JSON}）')
    args = parser.parse_args()

    data = load_data(args.o)

    if args.test:
        data = cli_mode(data, args.o)
        out_word = os.path.join(
            os.path.dirname(args.o), '..', 'output', 'project_info_test.docx')
        out_word = os.path.normpath(out_word)
        make_test_word(data, out_word)
        print(f'\n輸出 JSON：{os.path.normpath(args.o)}')
    elif args.cli:
        cli_mode(data, args.o)
    else:
        gui = ProjectInfoGUI(data, args.o)
        gui.run()


if __name__ == '__main__':
    main()
